#!/usr/bin/env python3
"""Build Chinese and English LumiaKey functional guides from JSON content."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


DEFAULT_CONTENT = Path("docs/manuals/manual_content.json")
DEFAULT_OUTPUT_DIR = Path("docs/manuals")
DEFAULT_SCREENSHOTS_DIR = Path("docs/manuals/screenshots")
LANGUAGES = ("zh", "en")
FONT_LATIN = "Calibri"
FONT_CHINESE = "Microsoft YaHei"
INK = RGBColor(0x00, 0x00, 0x00)
CAPTION_INK = RGBColor(0x55, 0x55, 0x55)
TABLE_BORDER = "D9D9D9"
TABLE_HEADER = "1F4E78"
TABLE_ALTERNATE = "DDEBF7"
CONTENT_WIDTH_IN = 6.65
COMPARISON_COLUMN_WIDTHS = (1.45, 2.35, 2.85)


def set_run_font(run, *, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CHINESE)


def set_style_font(style, *, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_CHINESE)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = TABLE_BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description or title)


def page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, result, end))


def configure_document(language: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    body_size = 9.5 if language == "zh" else 10
    normal = document.styles["Normal"]
    set_style_font(normal, size=body_size)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    title_style = document.styles["Title"]
    set_style_font(title_style, size=18, bold=True)
    title_style.paragraph_format.space_after = Pt(8)
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.keep_together = True
    title_style.paragraph_format.page_break_before = False
    title_style.paragraph_format.widow_control = True
    title_style.font.underline = False
    title_ppr = title_style._element.get_or_add_pPr()
    p_bdr = title_ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        title_ppr.remove(p_bdr)

    heading = document.styles["Heading 1"]
    set_style_font(heading, size=14, bold=True)
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("LumiaKey 1.3")
    set_run_font(header_run, size=9, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    page_number_field(footer)
    return document


def set_document_metadata(document: Document, content: dict[str, Any]) -> None:
    properties = document.core_properties
    properties.title = "LumiaKey Functional Guide"
    properties.subject = "LumiaKey functional guide"
    properties.author = "LumiaKey Project"
    date = str(content.get("date", ""))
    source_commit = str(content.get("sourceCommit", ""))
    metadata = []
    version = str(content.get("version", ""))
    if version:
        metadata.append(f"Version: {version}")
    if date:
        metadata.append(f"Content date: {date}")
    if source_commit:
        metadata.append(f"Source commit: {source_commit}")
    if metadata:
        properties.subject = "LumiaKey functional guide | " + " | ".join(metadata)


def localize(value: Any, language: str, default: str = "") -> str:
    if isinstance(value, dict):
        value = value.get(language, default)
    if value is None:
        return default
    return str(value)


def add_text_paragraph(document: Document, text: str, *, size: float, bold: bool = False,
                       keep_next: bool = False, alignment=None, after: float = 4):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = True
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_steps(document: Document, steps: list[dict[str, Any]], language: str, body_size: float) -> None:
    for index, step in enumerate(steps, start=1):
        text = localize(step, language)
        if not text:
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_together = True
        number = paragraph.add_run(f"{index}. ")
        set_run_font(number, size=body_size, bold=True)
        content = paragraph.add_run(text)
        set_run_font(content, size=body_size)


def add_comparison_table(document: Document, table_data: dict[str, Any], language: str,
                         body_size: float) -> None:
    headers = table_data.get("headers", {}).get(language, [])
    rows = [row.get(language, []) for row in table_data.get("rows", [])]
    if not headers or not rows:
        return
    column_count = len(headers)
    table = document.add_table(rows=1, cols=column_count)
    table.autofit = False
    widths = (COMPARISON_COLUMN_WIDTHS if column_count == 3 else
              tuple(CONTENT_WIDTH_IN / column_count for _ in range(column_count)))
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    set_table_borders(table)
    for col, value in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.width = Inches(widths[col])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, TABLE_HEADER)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        set_run_font(run, size=body_size, bold=True, color=RGBColor(255, 255, 255))
    prevent_row_split(table.rows[0])
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for col in range(column_count):
            cell = cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, TABLE_ALTERNATE if row_index % 2 == 0 else "FFFFFF")
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            value = row_values[col] if col < len(row_values) else ""
            set_run_font(paragraph.add_run(str(value)), size=body_size)
        prevent_row_split(table.rows[-1])


def image_size_inches(path: Path, max_width: float, max_height: float) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        raise ValueError(f"Invalid image dimensions: {path}")
    scale = min(max_width / width_px, max_height / height_px)
    return width_px * scale, height_px * scale


def add_image_item(paragraph, image_item: dict[str, Any], language: str,
                   screenshots_dir: Path, width_cap: float, height_cap: float) -> None:
    image_path = screenshots_dir / str(image_item["file"])
    if not image_path.is_file():
        raise FileNotFoundError(f"Screenshot not found: {image_path}")
    width, height = image_size_inches(image_path, width_cap, height_cap)
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    caption = localize(image_item.get("caption"), language)
    set_alt_text(shape, caption or image_path.stem, caption or image_path.name)


def add_screenshots(document: Document, images: list[dict[str, Any]], language: str,
                    screenshots_dir: Path, body_size: float, page: dict[str, Any]) -> None:
    if not images:
        return
    page_height_override = page.get("imageHeight")
    page_width_override = page.get("maxWidth")
    if len(images) == 1:
        item = images[0]
        width = float(item.get("maxWidth", page_width_override or CONTENT_WIDTH_IN))
        height = float(item.get("imageHeight", page_height_override or 2.9))
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if page.get("table"):
            paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        add_image_item(paragraph, item, language, screenshots_dir, min(width, CONTENT_WIDTH_IN), height)
        caption = localize(item.get("caption"), language)
        if caption:
            cap = add_text_paragraph(document, caption, size=8, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                     after=5)
            cap.paragraph_format.keep_together = True
            cap.paragraph_format.keep_with_next = False
            for run in cap.runs:
                run.font.color.rgb = CAPTION_INK
        return

    # Pair portrait screenshots in one paragraph with centered tab stops.
    # Landscape screenshots get their own row so controls remain legible.
    i = 0
    while i < len(images):
        first = images[i]
        first_path = screenshots_dir / str(first["file"])
        with Image.open(first_path) as image:
            first_portrait = image.height > image.width
        second = images[i + 1] if i + 1 < len(images) else None
        second_portrait = False
        if second is not None:
            second_path = screenshots_dir / str(second["file"])
            with Image.open(second_path) as image:
                second_portrait = image.height > image.width
        if second is not None and first_portrait and second_portrait:
            centers = (1.575, 4.725)
            image_p = document.add_paragraph()
            image_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if page.get("table") and i == 0:
                image_p.paragraph_format.space_before = Pt(6)
            image_p.paragraph_format.space_after = Pt(1)
            image_p.paragraph_format.keep_with_next = True
            image_p.paragraph_format.keep_together = True
            for center in centers:
                image_p.paragraph_format.tab_stops.add_tab_stop(Inches(center), WD_TAB_ALIGNMENT.CENTER)
            for item in (first, second):
                image_p.add_run("\t")
                width = float(item.get("maxWidth", 2.1))
                height = float(item.get("imageHeight", page_height_override or 4.65))
                add_image_item(image_p, item, language, screenshots_dir, min(width, 2.1), height)

            captions = [localize(item.get("caption"), language) for item in (first, second)]
            if any(captions):
                caption_p = document.add_paragraph()
                caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                caption_p.paragraph_format.space_after = Pt(4)
                caption_p.paragraph_format.keep_together = True
                for center in centers:
                    caption_p.paragraph_format.tab_stops.add_tab_stop(Inches(center), WD_TAB_ALIGNMENT.CENTER)
                for caption in captions:
                    caption_p.add_run("\t")
                    run = caption_p.add_run(caption)
                    set_run_font(run, size=8, color=CAPTION_INK)
            i += 2
        else:
            item = first
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if page.get("table") and i == 0:
                paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            width = float(item.get("maxWidth", page_width_override or CONTENT_WIDTH_IN))
            default_height = ((4.65 if first_portrait else 2.2)
                              if len(images) > 1 else 2.9)
            height = float(item.get("imageHeight", page_height_override or default_height))
            add_image_item(paragraph, item, language, screenshots_dir,
                           min(width, CONTENT_WIDTH_IN), height)
            caption = localize(item.get("caption"), language)
            if caption:
                cap = add_text_paragraph(document, caption, size=8,
                                         alignment=WD_ALIGN_PARAGRAPH.CENTER, after=4)
                for run in cap.runs:
                    run.font.color.rgb = CAPTION_INK
            i += 1


def add_page(document: Document, page: dict[str, Any], language: str,
             screenshots_dir: Path, is_first: bool) -> None:
    title = localize(page.get("title"), language)
    page_break_applied = False
    if title:
        title_p = document.add_paragraph(style="Title" if is_first else "Heading 1")
        if not is_first:
            title_p.paragraph_format.page_break_before = True
            page_break_applied = True
        title_p.paragraph_format.keep_with_next = True
        title_p.paragraph_format.keep_together = True
        run = title_p.add_run(title)
        set_run_font(run, size=18 if is_first else 14, bold=True)

    body_size = 9.5 if language == "zh" else 10
    intro = localize(page.get("intro"), language)
    if intro:
        intro_p = add_text_paragraph(document, intro, size=body_size)
        if not is_first and not page_break_applied:
            intro_p.paragraph_format.page_break_before = True
            page_break_applied = True

    steps = page.get("steps", [])
    if steps:
        add_steps(document, steps, language, body_size)

    for paragraph_text in page.get("paragraphs", []):
        text = localize(paragraph_text, language)
        if text:
            add_text_paragraph(document, text, size=body_size)

    table_data = page.get("table")
    if table_data:
        add_comparison_table(document, table_data, language, body_size)

    add_screenshots(document, page.get("images", []), language, screenshots_dir,
                    body_size, page)


def build_manual(content_path: Path, output_dir: Path, screenshots_dir: Path,
                 language: str) -> Path:
    content = json.loads(content_path.read_text(encoding="utf-8"))
    pages = content.get("pages")
    if not isinstance(pages, list) or not pages:
        raise ValueError("manual_content.json must contain a non-empty pages array")
    document = configure_document(language)
    set_document_metadata(document, content)
    for index, page in enumerate(pages):
        if not isinstance(page, dict):
            raise ValueError(f"pages[{index}] must be an object")
        add_page(document, page, language, screenshots_dir, index == 0)

    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = "ZH" if language == "zh" else "EN"
    output_path = output_dir / f"LumiaKey_Functional_Guide_{suffix}.docx"
    document.save(output_path)
    return output_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--content", type=Path, default=DEFAULT_CONTENT,
                        help="Path to bilingual manual content JSON")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR,
                        help="Directory for generated DOCX files")
    parser.add_argument("--screenshots-dir", type=Path, default=DEFAULT_SCREENSHOTS_DIR,
                        help="Directory containing screenshot image files")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        for language in LANGUAGES:
            output_path = build_manual(args.content, args.output_dir,
                                       args.screenshots_dir, language)
            print(output_path)
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
